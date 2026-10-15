"""ONS, "System Average Price (SAP) of gas", the workbook's sheet
`1.Daily SAP Gas`: one row per calendar day from 01/01/2018, column A a date
cell, column B `SAP actual day` in pence per kilowatt-hour (National Gas
Transmission's On-the-Day Commodity Market, a within-day system price).
Read as 024's schema pass found it (`archives/024-ons-sap-gas/`): header
on row 5, 3,206 dated rows to 11/10/2026, every `SAP actual day` cell a
number. Needs openpyxl (the `registers` dependency group).
"""

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

SHEET = "1.Daily SAP Gas"
HEADER_ROW = 5
HEADER = ("Date", "SAP actual day", "SAP seven-day rolling average")
PENCE_PER_KWH_TO_GBP_PER_MWH = Decimal(10)


def _rows(path: Path) -> tuple[list[str], list[tuple]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = [ws.title for ws in wb.worksheets]
    if SHEET not in sheets:
        raise ValueError(f"sheet {SHEET!r} not in {sheets}")
    return sheets, list(wb[SHEET].iter_rows(values_only=True))


def format_profile(path: Path) -> dict:
    """Sheet names, header cells and cell forms, with no value read."""
    sheets, rows = _rows(path)
    header = (
        [str(c) for c in rows[HEADER_ROW - 1] if c is not None] if len(rows) >= HEADER_ROW else []
    )
    body = [r for r in rows[HEADER_ROW:] if any(c is not None for c in r)]
    forms = {"date": {"date-cell": 0, "other": 0}, "sap": {"number": 0, "blank": 0, "other": 0}}
    for r in body:
        forms["date"]["date-cell" if isinstance(r[0], datetime | date) else "other"] += 1
        v = r[1] if len(r) > 1 else None
        forms["sap"][
            "blank"
            if v is None
            else "number"
            if isinstance(v, int | float) and not isinstance(v, bool)
            else "other"
        ] += 1
    return {
        "sheets": sheets,
        "header_row": HEADER_ROW,
        "header": header,
        "rows": len(body),
        "forms": forms,
    }


def as_decimal(v: object) -> Decimal:
    """A cell's number as the decimal its shortest repr spells: openpyxl
    hands back a float for `1.7083`, whose repr is `1.7083`."""
    if isinstance(v, bool) or not isinstance(v, int | float):
        raise ValueError(f"not a number cell: {v!r}")
    return Decimal(repr(v)) if isinstance(v, float) else Decimal(v)


def read(path: Path) -> dict[date, Decimal]:
    """Day -> SAP in GBP/MWh (the printed p/kWh times ten)."""
    _, rows = _rows(path)
    if tuple(str(c) for c in rows[HEADER_ROW - 1] if c is not None) != HEADER:
        raise ValueError(f"unexpected header {rows[HEADER_ROW - 1]!r}")
    out: dict[date, Decimal] = {}
    for r in rows[HEADER_ROW:]:
        if not any(c is not None for c in r):
            continue
        d = r[0]
        if not isinstance(d, datetime | date):
            raise ValueError(f"column A is not a date cell: {d!r}")
        day = d.date() if isinstance(d, datetime) else d
        if day in out:
            raise ValueError(f"duplicate date {day}")
        out[day] = as_decimal(r[1]) * PENCE_PER_KWH_TO_GBP_PER_MWH
    return out
